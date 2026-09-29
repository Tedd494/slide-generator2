"""
AI Pro Deck Generator
=====================
A single-file Streamlit application that turns a topic or raw notes into a
polished, multi-slide, professionally-themed presentation deck using 100%
FREE AI backends:

  - Text/outline generation: Groq (free API key, best quality) with automatic
    fallback to Pollinations.ai (zero setup, no key at all, genuinely free).
  - Image generation: Pollinations.ai's Flux model (free, no API key/signup
    required for every slide illustration).
  - Pillow (PIL) composites each slide (1920x1080) using one of several
    professional layout templates (title, content w/ image, big-stat,
    quote, closing) with theme colors, gradients, accent bars, logo
    placement, and automatic text wrapping / font auto-fit.
  - python-pptx exports the whole deck as an EDITABLE PowerPoint file
    (native text boxes + embedded images + gradient backgrounds), alongside
    a ZIP of flattened PNGs for quick sharing.

Run with:  streamlit run app.py
"""
import io
import json
import random
import re
import textwrap
import urllib.parse
import zipfile
from concurrent.futures import ThreadPoolExecutor, as_completed
from typing import Dict, List, Optional, Tuple

import requests
import streamlit as st
from PIL import Image, ImageDraw, ImageFont, ImageOps

try:
    from openai import OpenAI  # used ONLY as an OpenAI-compatible client for Groq
except ImportError:  # pragma: no cover
    OpenAI = None

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt, Emu


# --------------------------------------------------------------------------
# Constants & geometry
# --------------------------------------------------------------------------

CANVAS_W, CANVAS_H = 1920, 1080
MARGIN = 90
TOP_INSET = 150
BOTTOM_INSET = 110
PX_TO_IN = 13.333 / CANVAS_W  # keeps PPTX (16:9, 13.333in x 7.5in) geometry aligned with PNG px

FONT_DIRS_REGULAR = [
    "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
    "/usr/share/fonts/truetype/liberation/LiberationSans-Regular.ttf",
    "DejaVuSans.ttf",
]
FONT_DIRS_BOLD = [
    "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf",
    "/usr/share/fonts/truetype/liberation/LiberationSans-Bold.ttf",
    "DejaVuSans-Bold.ttf",
]
FONT_DIRS_ITALIC = [
    "/usr/share/fonts/truetype/dejavu/DejaVuSans-Oblique.ttf",
    "/usr/share/fonts/truetype/liberation/LiberationSans-Italic.ttf",
    "DejaVuSans-Oblique.ttf",
]

LOGO_POSITIONS = ["Top-Left", "Top-Right", "Bottom-Left", "Bottom-Right"]

THEMES = {
    "Midnight Blue":     {"bg1": (12, 17, 38),  "bg2": (24, 34, 74),  "accent": "#5B8CFF", "text": "#FFFFFF", "muted": "#AEB9E8", "light": False},
    "Charcoal & Gold":   {"bg1": (23, 21, 18),  "bg2": (41, 36, 27),  "accent": "#D4AF37", "text": "#F6F1E6", "muted": "#CBBE9C", "light": False},
    "Emerald Executive": {"bg1": (8, 26, 21),   "bg2": (14, 42, 34),  "accent": "#33D48A", "text": "#EFFBF5", "muted": "#A9E3C6", "light": False},
    "Crimson Bold":      {"bg1": (28, 11, 15),  "bg2": (48, 18, 22),  "accent": "#FF5A4E", "text": "#FFF4F2", "muted": "#F0B3AC", "light": False},
    "Slate Minimal":     {"bg1": (246, 247, 250), "bg2": (231, 234, 240), "accent": "#2563EB", "text": "#12151C", "muted": "#5B6472", "light": True},
}

LAYOUT_CHOICES = [
    "title", "content_image_right", "content_image_left",
    "content_full", "big_stat", "quote", "closing",
]


# --------------------------------------------------------------------------
# Free AI text backend: Groq (optional key) -> Pollinations.ai (no key, always works)
# --------------------------------------------------------------------------

def call_free_llm(messages: List[dict], groq_key: str = "") -> str:
    """
    Returns the assistant's raw text content.
    Tries Groq first (if a free key is supplied), then falls back to the
    completely free, no-signup Pollinations.ai text API.
    """
    if groq_key and OpenAI is not None:
        try:
            client = OpenAI(base_url="https://api.groq.com/openai/v1", api_key=groq_key)
            resp = client.chat.completions.create(
                model="llama-3.3-70b-versatile",
                messages=messages,
                temperature=0.6,
                max_tokens=1800,
            )
            return resp.choices[0].message.content
        except Exception as e:
            st.warning(f"Groq call failed ({e}); falling back to Pollinations.ai (free, no key).")

    # Pollinations.ai - zero-auth, zero-signup free text API
    r = requests.post(
        "https://text.pollinations.ai/openai",
        json={
            "messages": messages,
            "model": "openai",
            "private": True,
            "seed": random.randint(1, 999_999_999),
        },
        timeout=90,
    )
    r.raise_for_status()
    return r.json()["choices"][0]["message"]["content"]


def extract_json(text: str) -> dict:
    """Robustly pull a JSON object out of an LLM response, tolerating stray prose/fences."""
    text = text.strip()
    text = re.sub(r"^```(json)?", "", text.strip(), flags=re.IGNORECASE).strip()
    text = re.sub(r"```$", "", text.strip()).strip()
    try:
        return json.loads(text)
    except Exception:
        start = text.find("{")
        end = text.rfind("}")
        if start != -1 and end != -1 and end > start:
            return json.loads(text[start:end + 1])
        raise


PLANNER_SYSTEM = """You are an expert presentation designer. Output ONLY raw JSON \
(no markdown fences, no commentary) matching this schema exactly:

{
  "deck_title": "string, short punchy deck title",
  "slides": [
    {"layout": "title", "title": "string", "subtitle": "string"},
    {"layout": "content_image_right", "title": "string", "bullets": ["string", "string", "string"], "image_prompt": "string, vivid visual description, no text/words in image"},
    {"layout": "content_image_left", "title": "string", "bullets": ["string", "string", "string"], "image_prompt": "string"},
    {"layout": "content_full", "title": "string", "bullets": ["string", "string", "string", "string"]},
    {"layout": "big_stat", "stat": "short number or percentage e.g. '+24%'", "caption": "one short supporting sentence"},
    {"layout": "quote", "quote": "short inspiring or relevant quote", "author": "attribution"},
    {"layout": "closing", "title": "string e.g. Thank You", "subtitle": "string, call to action or contact info"}
  ]
}

Rules:
- Produce EXACTLY the number of slides requested by the user.
- The FIRST slide must always be layout "title" and the LAST slide must always be layout "closing".
- Alternate between content_image_right / content_image_left for visual variety.
- Include at most one "big_stat" and one "quote" slide, only if it fits the topic naturally.
- Bullets must be short (under 14 words each), plain text, no markdown, no numbering.
- image_prompt must describe a clean, modern, professional illustration (flat design or 3D render style), and must NOT ask for any text/words/letters to appear in the image.
"""


def generate_deck_plan(topic: str, num_slides: int, groq_key: str = "") -> dict:
    user_msg = (
        f"Create a {num_slides}-slide presentation deck about:\n\n{topic}\n\n"
        f"Return exactly {num_slides} slide objects total (including the title and closing slides)."
    )
    raw = call_free_llm(
        [{"role": "system", "content": PLANNER_SYSTEM}, {"role": "user", "content": user_msg}],
        groq_key=groq_key,
    )
    plan = extract_json(raw)
    if "slides" not in plan or not isinstance(plan["slides"], list) or not plan["slides"]:
        raise ValueError("AI did not return a usable slide plan. Try again or rephrase your topic.")
    return plan


# --------------------------------------------------------------------------
# Free AI image backend: Pollinations.ai (Flux model, no key required)
# --------------------------------------------------------------------------

def generate_pollinations_image(prompt: str, width: int = 1024, height: int = 1024) -> Image.Image:
    encoded = urllib.parse.quote(prompt)
    seed = random.randint(1, 999_999_999)
    url = (
        f"https://image.pollinations.ai/prompt/{encoded}"
        f"?width={width}&height={height}&nologo=true&model=flux&seed={seed}&safe=true"
    )
    r = requests.get(url, timeout=120)
    r.raise_for_status()
    return Image.open(io.BytesIO(r.content)).convert("RGBA")


def generate_images_for_plan(plan: dict) -> Dict[int, Image.Image]:
    """Generate all slide illustrations concurrently (Pollinations requests are independent HTTP GETs)."""
    jobs = {}
    for i, slide in enumerate(plan["slides"]):
        prompt = slide.get("image_prompt")
        if prompt and slide.get("layout") in ("content_image_right", "content_image_left"):
            jobs[i] = prompt

    results: Dict[int, Image.Image] = {}
    if not jobs:
        return results

    with ThreadPoolExecutor(max_workers=min(6, len(jobs))) as pool:
        future_map = {pool.submit(generate_pollinations_image, p, 1024, 1024): i for i, p in jobs.items()}
        for future in as_completed(future_map):
            idx = future_map[future]
            try:
                results[idx] = future.result()
            except Exception as e:
                st.warning(f"Image generation failed for slide {idx + 1}: {e}")
    return results


# --------------------------------------------------------------------------
# Pillow helpers
# --------------------------------------------------------------------------

def load_font(size: int, bold: bool = False, italic: bool = False) -> ImageFont.FreeTypeFont:
    if italic:
        candidates = FONT_DIRS_ITALIC
    elif bold:
        candidates = FONT_DIRS_BOLD
    else:
        candidates = FONT_DIRS_REGULAR
    for path in candidates:
        try:
            return ImageFont.truetype(path, size)
        except Exception:
            continue
    try:
        return ImageFont.load_default(size=size)
    except TypeError:
        return ImageFont.load_default()


def wrap_text_to_width(draw: ImageDraw.ImageDraw, text: str, font: ImageFont.FreeTypeFont, max_width: int) -> List[str]:
    words = text.split()
    if not words:
        return [""]
    lines, current = [], ""
    for word in words:
        candidate = f"{current} {word}".strip()
        if draw.textlength(candidate, font=font) <= max_width or not current:
            current = candidate
        else:
            lines.append(current)
            current = word
    if current:
        lines.append(current)
    return lines


def make_gradient_bg(size: Tuple[int, int], color1: Tuple[int, int, int], color2: Tuple[int, int, int],
                      direction: str = "diagonal") -> Image.Image:
    w, h = size
    base = Image.linear_gradient("L")
    if direction == "horizontal":
        base = base.rotate(90, expand=True).resize((w, h))
    elif direction == "diagonal":
        s = max(w, h) * 2
        big = base.resize((s, s)).rotate(45, expand=True)
        bw, bh = big.size
        left, top = (bw - w) // 2, (bh - h) // 2
        base = big.crop((left, top, left + w, top + h))
    else:
        base = base.resize((w, h))
    return ImageOps.colorize(base.convert("L"), black=color1, white=color2).convert("RGB")


def hex_to_rgb(hex_color: str) -> Tuple[int, int, int]:
    hex_color = hex_color.lstrip("#")
    return tuple(int(hex_color[i:i + 2], 16) for i in (0, 2, 4))


def paste_logo(canvas: Image.Image, logo: Optional[Image.Image], position: str) -> None:
    if logo is None:
        return
    logo = logo.convert("RGBA").copy()
    max_w, max_h = int(CANVAS_W * 0.11), int(CANVAS_H * 0.11)
    logo.thumbnail((max_w, max_h), Image.LANCZOS)
    lw, lh = logo.size
    margin = 46
    coords = {
        "Top-Left": (margin, margin),
        "Top-Right": (CANVAS_W - lw - margin, margin),
        "Bottom-Left": (margin, CANVAS_H - lh - margin),
        "Bottom-Right": (CANVAS_W - lw - margin, CANVAS_H - lh - margin),
    }
    x, y = coords.get(position, coords["Top-Right"])
    canvas.paste(logo, (x, y), logo)


def draw_footer(draw: ImageDraw.ImageDraw, theme: dict, deck_title: str, slide_no: int, total: int) -> None:
    font = load_font(22)
    muted = theme["muted"]
    label = f"{deck_title}"
    draw.text((MARGIN, CANVAS_H - 56), label, font=font, fill=muted)
    counter = f"{slide_no:02d} / {total:02d}"
    w = draw.textlength(counter, font=font)
    draw.text((CANVAS_W - MARGIN - w, CANVAS_H - 56), counter, font=font, fill=muted)


def draw_bullets(draw: ImageDraw.ImageDraw, bullets: List[str], x: int, y: int, box_width: int,
                  font: ImageFont.FreeTypeFont, color: str, accent_hex: str,
                  line_height: float, gap: float) -> int:
    indent = 42
    marker_r = 7
    for bullet in bullets:
        lines = wrap_text_to_width(draw, bullet, font, box_width - indent)
        marker_cy = y + line_height * 0.38
        draw.ellipse([x, marker_cy - marker_r, x + marker_r * 2, marker_cy + marker_r], fill=accent_hex)
        for line in lines:
            draw.text((x + indent, y), line, font=font, fill=color)
            y += line_height
        y += gap
    return y


def measure_bullets_height(draw, bullets, font, box_width, line_height, gap) -> float:
    indent = 42
    total = 0.0
    for bullet in bullets:
        lines = wrap_text_to_width(draw, bullet, font, box_width - indent)
        total += len(lines) * line_height + gap
    return total


def fit_content_font(draw, title: str, bullets: List[str], box_w: int, box_h: int,
                      base_size: int, min_size: int = 18):
    size = base_size
    while size >= min_size:
        title_font = load_font(size + 12, bold=True)
        body_font = load_font(size)
        title_lines = wrap_text_to_width(draw, title, title_font, box_w)
        title_h = len(title_lines) * (size + 12) * 1.25 + 40
        line_h = size * 1.5
        gap = size * 0.55
        bullets_h = measure_bullets_height(draw, bullets, body_font, box_w, line_h, gap)
        if title_h + bullets_h <= box_h or size <= min_size:
            return title_font, body_font, title_lines, title_h, line_h, gap
        size -= 2
    return title_font, body_font, title_lines, title_h, line_h, gap


# --------------------------------------------------------------------------
# Slide renderers (each returns a fresh 1920x1080 RGB PIL.Image)
# --------------------------------------------------------------------------

def render_title_slide(slide: dict, theme: dict, accent: str, logo, deck_title: str, slide_no: int, total: int) -> Image.Image:
    canvas = make_gradient_bg((CANVAS_W, CANVAS_H), theme["bg1"], theme["bg2"], "diagonal")
    draw = ImageDraw.Draw(canvas)

    # decorative accent ring, top-left, for visual interest
    draw.ellipse([-180, -180, 260, 260], outline=accent, width=8)
    draw.ellipse([CANVAS_W - 260, CANVAS_H - 260, CANVAS_W + 180, CANVAS_H + 180], outline=accent, width=8)

    title = slide.get("title", deck_title)
    subtitle = slide.get("subtitle", "")

    title_font = load_font(96, bold=True)
    box_w = CANVAS_W - 2 * (MARGIN + 120)
    title_lines = wrap_text_to_width(draw, title, title_font, box_w)
    sub_font = load_font(38)
    sub_lines = wrap_text_to_width(draw, subtitle, sub_font, box_w) if subtitle else []

    total_h = len(title_lines) * 112 + (40 if sub_lines else 0) + len(sub_lines) * 52 + 40
    y = (CANVAS_H - total_h) // 2

    for line in title_lines:
        w = draw.textlength(line, font=title_font)
        draw.text(((CANVAS_W - w) // 2, y), line, font=title_font, fill=theme["text"])
        y += 112

    # accent underline
    line_w = 160
    draw.rectangle([(CANVAS_W - line_w) // 2, y + 6, (CANVAS_W + line_w) // 2, y + 14], fill=accent)
    y += 46

    for line in sub_lines:
        w = draw.textlength(line, font=sub_font)
        draw.text(((CANVAS_W - w) // 2, y), line, font=sub_font, fill=theme["muted"])
        y += 52

    paste_logo(canvas, logo, "Top-Right" if logo else "Top-Right")
    return canvas


def render_content_slide(slide: dict, theme: dict, accent: str, logo, image: Optional[Image.Image],
                          image_side: str, deck_title: str, slide_no: int, total: int) -> Image.Image:
    canvas = Image.new("RGB", (CANVAS_W, CANVAS_H), theme["bg1"])
    # subtle vertical vignette for depth without hurting text contrast
    grad = make_gradient_bg((CANVAS_W, CANVAS_H), theme["bg1"], theme["bg2"], "vertical")
    canvas = Image.blend(canvas, grad, 0.35)
    draw = ImageDraw.Draw(canvas)

    # left accent bar
    draw.rectangle([0, 0, 10, CANVAS_H], fill=accent)

    half = CANVAS_W // 2
    gutter = 50
    if image_side == "right":
        text_rect = (MARGIN, TOP_INSET, half - gutter, CANVAS_H - BOTTOM_INSET)
        image_rect = (half + gutter, TOP_INSET, CANVAS_W - MARGIN, CANVAS_H - BOTTOM_INSET)
    elif image_side == "left":
        image_rect = (MARGIN, TOP_INSET, half - gutter, CANVAS_H - BOTTOM_INSET)
        text_rect = (half + gutter, TOP_INSET, CANVAS_W - MARGIN, CANVAS_H - BOTTOM_INSET)
    else:
        text_rect = (MARGIN, TOP_INSET, CANVAS_W - MARGIN, CANVAS_H - BOTTOM_INSET)
        image_rect = None

    title = slide.get("title", "")
    bullets = slide.get("bullets", [])
    x0, y0, x1, y1 = text_rect
    box_w, box_h = x1 - x0, y1 - y0

    title_font, body_font, title_lines, title_h, line_h, gap = fit_content_font(
        draw, title, bullets, box_w, box_h, base_size=40
    )

    y = y0
    for line in title_lines:
        draw.text((x0, y), line, font=title_font, fill=theme["text"])
        y += title_font.size * 1.25
    y += 18
    draw.rectangle([x0, y, x0 + 90, y + 6], fill=accent)
    y += 34

    draw_bullets(draw, bullets, x0, y, box_w, body_font, theme["text"], accent, line_h, gap)

    if image_rect is not None and image is not None:
        ix0, iy0, ix1, iy1 = image_rect
        box_w2, box_h2 = ix1 - ix0, iy1 - iy0
        img = image.convert("RGBA").copy()
        fitted = ImageOps.fit(img, (box_w2, box_h2), method=Image.LANCZOS)
        # rounded-rect mask for a polished framed look
        mask = Image.new("L", fitted.size, 0)
        ImageDraw.Draw(mask).rounded_rectangle([0, 0, fitted.size[0], fitted.size[1]], radius=28, fill=255)
        # accent-colored border behind the frame
        border_pad = 10
        border_box = Image.new("RGBA", canvas.size, (0, 0, 0, 0))
        ImageDraw.Draw(border_box).rounded_rectangle(
            [ix0 - border_pad, iy0 - border_pad, ix1 + border_pad, iy1 + border_pad],
            radius=32, outline=accent, width=4
        )
        canvas.paste(Image.alpha_composite(canvas.convert("RGBA"), border_box).convert("RGB"), (0, 0))
        canvas.paste(fitted, (ix0, iy0), mask)

    draw = ImageDraw.Draw(canvas)
    paste_logo(canvas, logo, "Top-Right")
    draw_footer(draw, theme, deck_title, slide_no, total)
    return canvas


def render_big_stat_slide(slide: dict, theme: dict, accent: str, logo, deck_title: str, slide_no: int, total: int) -> Image.Image:
    canvas = make_gradient_bg((CANVAS_W, CANVAS_H), theme["bg1"], theme["bg2"], "diagonal")
    draw = ImageDraw.Draw(canvas)

    stat = slide.get("stat", "")
    caption = slide.get("caption", "")

    # decorative ring behind the stat
    ring_r = 300
    cx, cy = CANVAS_W // 2, CANVAS_H // 2 - 40
    draw.ellipse([cx - ring_r, cy - ring_r, cx + ring_r, cy + ring_r], outline=accent, width=10)

    stat_font = load_font(160, bold=True)
    w = draw.textlength(stat, font=stat_font)
    draw.text((cx - w / 2, cy - 95), stat, font=stat_font, fill=theme["text"])

    cap_font = load_font(36)
    box_w = CANVAS_W - 2 * (MARGIN + 260)
    cap_lines = wrap_text_to_width(draw, caption, cap_font, box_w)
    cy2 = cy + ring_r + 60
    for line in cap_lines:
        w = draw.textlength(line, font=cap_font)
        draw.text(((CANVAS_W - w) // 2, cy2), line, font=cap_font, fill=theme["muted"])
        cy2 += 48

    paste_logo(canvas, logo, "Top-Right")
    draw_footer(draw, theme, deck_title, slide_no, total)
    return canvas


def render_quote_slide(slide: dict, theme: dict, accent: str, logo, deck_title: str, slide_no: int, total: int) -> Image.Image:
    canvas = make_gradient_bg((CANVAS_W, CANVAS_H), theme["bg2"], theme["bg1"], "vertical")
    draw = ImageDraw.Draw(canvas)

    mark_font = load_font(220, bold=True)
    draw.text((CANVAS_W // 2 - 70, 100), "\u201C", font=mark_font, fill=accent)

    quote = slide.get("quote", "")
    author = slide.get("author", "")

    quote_font = load_font(56, italic=True)
    box_w = CANVAS_W - 2 * (MARGIN + 220)
    lines = wrap_text_to_width(draw, quote, quote_font, box_w)
    total_h = len(lines) * 74
    y = (CANVAS_H - total_h) // 2
    for line in lines:
        w = draw.textlength(line, font=quote_font)
        draw.text(((CANVAS_W - w) // 2, y), line, font=quote_font, fill=theme["text"])
        y += 74

    y += 30
    line_w = 100
    draw.rectangle([(CANVAS_W - line_w) // 2, y, (CANVAS_W + line_w) // 2, y + 6], fill=accent)
    y += 30
    if author:
        author_font = load_font(32, bold=True)
        w = draw.textlength(f"— {author}", font=author_font)
        draw.text(((CANVAS_W - w) // 2, y), f"— {author}", font=author_font, fill=theme["muted"])

    paste_logo(canvas, logo, "Top-Right")
    draw_footer(draw, theme, deck_title, slide_no, total)
    return canvas


def render_closing_slide(slide: dict, theme: dict, accent: str, logo, deck_title: str, slide_no: int, total: int) -> Image.Image:
    canvas = make_gradient_bg((CANVAS_W, CANVAS_H), theme["bg1"], theme["bg2"], "diagonal")
    draw = ImageDraw.Draw(canvas)

    title = slide.get("title", "Thank You")
    subtitle = slide.get("subtitle", "")

    title_font = load_font(110, bold=True)
    w = draw.textlength(title, font=title_font)
    y = CANVAS_H // 2 - 90
    draw.text(((CANVAS_W - w) // 2, y), title, font=title_font, fill=theme["text"])
    y += 128
    line_w = 140
    draw.rectangle([(CANVAS_W - line_w) // 2, y, (CANVAS_W + line_w) // 2, y + 6], fill=accent)
    y += 40

    if subtitle:
        sub_font = load_font(36)
        box_w = CANVAS_W - 2 * (MARGIN + 260)
        for line in wrap_text_to_width(draw, subtitle, sub_font, box_w):
            w = draw.textlength(line, font=sub_font)
            draw.text(((CANVAS_W - w) // 2, y), line, font=sub_font, fill=theme["muted"])
            y += 48

    paste_logo(canvas, logo, "Top-Right")
    return canvas


def render_slide(slide: dict, idx: int, total: int, deck_title: str, theme: dict, accent: str,
                  logo, image: Optional[Image.Image]) -> Image.Image:
    layout = slide.get("layout", "content_full")
    slide_no = idx + 1
    if layout == "title":
        return render_title_slide(slide, theme, accent, logo, deck_title, slide_no, total)
    if layout == "content_image_right":
        return render_content_slide(slide, theme, accent, logo, image, "right", deck_title, slide_no, total)
    if layout == "content_image_left":
        return render_content_slide(slide, theme, accent, logo, image, "left", deck_title, slide_no, total)
    if layout == "content_full":
        return render_content_slide(slide, theme, accent, logo, None, "none", deck_title, slide_no, total)
    if layout == "big_stat":
        return render_big_stat_slide(slide, theme, accent, logo, deck_title, slide_no, total)
    if layout == "quote":
        return render_quote_slide(slide, theme, accent, logo, deck_title, slide_no, total)
    if layout == "closing":
        return render_closing_slide(slide, theme, accent, logo, deck_title, slide_no, total)
    return render_content_slide(slide, theme, accent, logo, image, "right", deck_title, slide_no, total)


# --------------------------------------------------------------------------
# PPTX export (native editable text boxes + embedded images + gradients)
# --------------------------------------------------------------------------

def _px_in(px: float) -> float:
    return px * PX_TO_IN


def _set_gradient_bg(pptx_slide, c1: Tuple[int, int, int], c2: Tuple[int, int, int], angle: float = 45.0):
    fill = pptx_slide.background.fill
    fill.gradient()
    stops = fill.gradient_stops
    stops[0].color.rgb = RGBColor(*c1)
    stops[1].color.rgb = RGBColor(*c2)
    try:
        fill.gradient_angle = angle
    except Exception:
        pass


def _add_textbox(pptx_slide, x, y, w, h, text, size, color_hex, bold=False, italic=False, align=PP_ALIGN.LEFT):
    box = pptx_slide.shapes.add_textbox(Inches(_px_in(x)), Inches(_px_in(y)), Inches(_px_in(w)), Inches(_px_in(h)))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = RGBColor(*hex_to_rgb(color_hex))
    return box


def _add_rect(pptx_slide, x, y, w, h, color_hex):
    from pptx.enum.shapes import MSO_SHAPE
    shp = pptx_slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(_px_in(x)), Inches(_px_in(y)), Inches(_px_in(w)), Inches(_px_in(h)))
    shp.fill.solid()
    shp.fill.fore_color.rgb = RGBColor(*hex_to_rgb(color_hex))
    shp.line.fill.background()
    return shp


def _add_bullets(pptx_slide, x, y, w, h, bullets, size, color_hex, accent_hex):
    box = pptx_slide.shapes.add_textbox(Inches(_px_in(x)), Inches(_px_in(y)), Inches(_px_in(w)), Inches(_px_in(h)))
    tf = box.text_frame
    tf.word_wrap = True
    for i, bullet in enumerate(bullets):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        run = p.add_run()
        run.text = f"\u25CF  {bullet}"
        run.font.size = Pt(size)
        run.font.color.rgb = RGBColor(*hex_to_rgb(color_hex))
        p.space_after = Pt(14)
    return box


def _add_picture_bytes(pptx_slide, pil_image: Image.Image, x, y, w, h):
    buf = io.BytesIO()
    pil_image.convert("RGB").save(buf, format="PNG")
    buf.seek(0)
    pptx_slide.shapes.add_picture(buf, Inches(_px_in(x)), Inches(_px_in(y)), Inches(_px_in(w)), Inches(_px_in(h)))


def build_pptx(plan: dict, theme: dict, accent: str, logo: Optional[Image.Image],
                images: Dict[int, Image.Image], logo_position: str) -> bytes:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    deck_title = plan.get("deck_title", "Untitled Deck")
    slides = plan["slides"]
    total = len(slides)
    text_hex = "#{:02X}{:02X}{:02X}".format(*hex_to_rgb(theme["text"])) if isinstance(theme["text"], tuple) else theme["text"]

    for idx, slide in enumerate(slides):
        layout = slide.get("layout", "content_full")
        s = prs.slides.add_slide(blank_layout)
        _set_gradient_bg(s, theme["bg1"], theme["bg2"], 45.0)

        if logo is not None:
            lw = CANVAS_W * 0.11
            lh = lw * (logo.height / logo.width)
            margin = 46
            coords = {
                "Top-Left": (margin, margin),
                "Top-Right": (CANVAS_W - lw - margin, margin),
                "Bottom-Left": (margin, CANVAS_H - lh - margin),
                "Bottom-Right": (CANVAS_W - lw - margin, CANVAS_H - lh - margin),
            }
            lx, ly = coords.get(logo_position, coords["Top-Right"])
            _add_picture_bytes(s, logo, lx, ly, lw, lh)

        if layout == "title":
            _add_textbox(s, 160, 380, CANVAS_W - 320, 160, slide.get("title", deck_title),
                         48, theme["text"], bold=True, align=PP_ALIGN.CENTER)
            _add_textbox(s, 160, 540, CANVAS_W - 320, 80, slide.get("subtitle", ""),
                         22, theme["muted"], align=PP_ALIGN.CENTER)

        elif layout in ("content_image_right", "content_image_left", "content_full"):
            half = CANVAS_W // 2
            gutter = 50
            if layout == "content_image_right":
                tx0, tx1 = MARGIN, half - gutter
                ix0, ix1 = half + gutter, CANVAS_W - MARGIN
            elif layout == "content_image_left":
                ix0, ix1 = MARGIN, half - gutter
                tx0, tx1 = half + gutter, CANVAS_W - MARGIN
            else:
                tx0, tx1 = MARGIN, CANVAS_W - MARGIN
                ix0 = ix1 = None

            _add_rect(s, tx0, TOP_INSET - 10, 6, CANVAS_H - TOP_INSET - BOTTOM_INSET + 10, accent)
            _add_textbox(s, tx0 + 30, TOP_INSET, tx1 - tx0 - 30, 100, slide.get("title", ""),
                         30, theme["text"], bold=True)
            _add_bullets(s, tx0 + 30, TOP_INSET + 110, tx1 - tx0 - 30, CANVAS_H - TOP_INSET - BOTTOM_INSET - 110,
                         slide.get("bullets", []), 16, theme["text"], accent)

            img = images.get(idx)
            if img is not None and ix0 is not None:
                _add_picture_bytes(s, img, ix0, TOP_INSET, ix1 - ix0, CANVAS_H - TOP_INSET - BOTTOM_INSET)

        elif layout == "big_stat":
            _add_textbox(s, 160, 380, CANVAS_W - 320, 160, slide.get("stat", ""),
                         72, theme["text"], bold=True, align=PP_ALIGN.CENTER)
            _add_textbox(s, 260, 560, CANVAS_W - 520, 100, slide.get("caption", ""),
                         20, theme["muted"], align=PP_ALIGN.CENTER)

        elif layout == "quote":
            _add_textbox(s, 260, 340, CANVAS_W - 520, 220, f"\u201C{slide.get('quote', '')}\u201D",
                         28, theme["text"], italic=True, align=PP_ALIGN.CENTER)
            _add_textbox(s, 260, 560, CANVAS_W - 520, 80, f"\u2014 {slide.get('author', '')}",
                         18, theme["muted"], bold=True, align=PP_ALIGN.CENTER)

        elif layout == "closing":
            _add_textbox(s, 160, 420, CANVAS_W - 320, 140, slide.get("title", "Thank You"),
                         48, theme["text"], bold=True, align=PP_ALIGN.CENTER)
            _add_textbox(s, 160, 560, CANVAS_W - 320, 80, slide.get("subtitle", ""),
                         20, theme["muted"], align=PP_ALIGN.CENTER)

        if layout not in ("title", "closing"):
            _add_textbox(s, MARGIN, CANVAS_H - 70, 600, 40, deck_title, 11, theme["muted"])
            _add_textbox(s, CANVAS_W - MARGIN - 200, CANVAS_H - 70, 200, 40,
                         f"{idx + 1:02d} / {total:02d}", 11, theme["muted"], align=PP_ALIGN.RIGHT)

    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


# --------------------------------------------------------------------------
# Streamlit App
# --------------------------------------------------------------------------

def init_state():
    defaults = {"deck_plan": None, "deck_images": {}, "deck_slides": None}
    for k, v in defaults.items():
        if k not in st.session_state:
            st.session_state[k] = v


def slide_editor(slide: dict, idx: int) -> dict:
    layout = st.selectbox("Layout", LAYOUT_CHOICES, index=LAYOUT_CHOICES.index(slide.get("layout", "content_full")),
                           key=f"layout_{idx}")
    slide["layout"] = layout

    if layout in ("title", "closing"):
        slide["title"] = st.text_input("Title", slide.get("title", ""), key=f"title_{idx}")
        slide["subtitle"] = st.text_input("Subtitle", slide.get("subtitle", ""), key=f"subtitle_{idx}")
    elif layout in ("content_image_right", "content_image_left", "content_full"):
        slide["title"] = st.text_input("Title", slide.get("title", ""), key=f"title_{idx}")
        bullets_text = st.text_area("Bullets (one per line)", "\n".join(slide.get("bullets", [])), key=f"bullets_{idx}", height=120)
        slide["bullets"] = [b.strip() for b in bullets_text.splitlines() if b.strip()]
        if layout != "content_full":
            slide["image_prompt"] = st.text_input("Image prompt", slide.get("image_prompt", ""), key=f"imgprompt_{idx}")
    elif layout == "big_stat":
        slide["stat"] = st.text_input("Stat (e.g. +24%)", slide.get("stat", ""), key=f"stat_{idx}")
        slide["caption"] = st.text_input("Caption", slide.get("caption", ""), key=f"caption_{idx}")
    elif layout == "quote":
        slide["quote"] = st.text_area("Quote", slide.get("quote", ""), key=f"quote_{idx}", height=80)
        slide["author"] = st.text_input("Author", slide.get("author", ""), key=f"author_{idx}")

    return slide


def main():
    st.set_page_config(page_title="AI Pro Deck Generator", layout="wide")
    init_state()

    st.title("🖼️ AI Pro Deck Generator — 100% Free API")
    st.caption("Groq (free key) or Pollinations.ai (zero setup) for text · Pollinations.ai Flux for images · Pillow for pro layouts · python-pptx for an editable deck.")

    with st.sidebar:
        st.header("🔌 Free AI Backend")
        groq_key = st.text_input(
            "Groq API Key (optional)", type="password",
            help="Free at console.groq.com/keys, no credit card. Leave blank to auto-use Pollinations.ai (also free, zero setup)."
        )
        if not groq_key:
            st.info("No Groq key set → using **Pollinations.ai** automatically (completely free, no signup needed) for both text and images.")
        else:
            st.success("Using Groq (Llama 3.3 70B) for text. Images always use free Pollinations.ai.")

        include_images = st.checkbox("Generate AI illustrations for content slides", value=True)

        st.divider()
        st.header("🎨 Brand & Theme")
        logo_file = st.file_uploader("Company Logo", type=["png", "jpg", "jpeg", "webp"])
        logo_image = Image.open(logo_file) if logo_file else None
        if logo_image:
            st.image(logo_image, caption="Logo preview", width=110)

        theme_name = st.selectbox("Theme", list(THEMES.keys()), index=0)
        theme = dict(THEMES[theme_name])
        accent = st.color_picker("Accent Color", value=theme["accent"])
        logo_position = st.selectbox("Logo Position", LOGO_POSITIONS, index=1)

        st.divider()
        st.header("📐 Deck Settings")
        num_slides = st.slider("Number of slides", min_value=4, max_value=10, value=6)

    st.subheader("1. Describe Your Presentation")
    topic = st.text_area(
        "Topic or raw content", height=150,
        placeholder="e.g. Our Q3 2026 sales results and Q4 strategy for a SaaS startup called Nimbus CRM. "
                    "Revenue grew 24%, churn dropped, launching AI features in Q4...",
    )
    if st.button("🧠 Generate Deck Outline with AI", type="primary", use_container_width=True):
        if not topic.strip():
            st.warning("Please describe your topic or paste some raw content first.")
        else:
            with st.spinner("Planning your deck with AI (free backend)..."):
                try:
                    st.session_state.deck_plan = generate_deck_plan(topic, num_slides, groq_key)
                    st.session_state.deck_images = {}
                    st.session_state.deck_slides = None
                    st.success(f"Outline ready: {st.session_state.deck_plan.get('deck_title', '')}")
                except Exception as e:
                    st.error(f"Could not generate an outline: {e}")

    plan = st.session_state.deck_plan
    if plan:
        st.divider()
        st.subheader("2. Review & Edit Slides")
        plan["deck_title"] = st.text_input("Deck Title", plan.get("deck_title", ""))
        for i, slide in enumerate(plan["slides"]):
            with st.expander(f"Slide {i + 1} — {slide.get('layout', '')}", expanded=False):
                plan["slides"][i] = slide_editor(slide, i)

        st.divider()
        st.subheader("3. Generate Illustrations & Compose Slides")
        if st.button("🖌️ Build the Deck", type="primary", use_container_width=True):
            with st.spinner("Generating free AI illustrations and compositing slides..."):
                images = generate_images_for_plan(plan) if include_images else {}
                st.session_state.deck_images = images

                deck_title = plan.get("deck_title", "Untitled Deck")
                total = len(plan["slides"])
                slides_rendered = []
                for i, slide in enumerate(plan["slides"]):
                    img = render_slide(slide, i, total, deck_title, theme, accent, logo_image, images.get(i))
                    slides_rendered.append(img)
                st.session_state.deck_slides = slides_rendered
            st.success("Deck composed! Scroll down to preview and download.")

    slides_rendered = st.session_state.deck_slides
    if slides_rendered:
        st.divider()
        st.subheader("4. Preview & Download")
        idx = st.slider("Preview slide", 1, len(slides_rendered), 1) - 1
        st.image(slides_rendered[idx], use_container_width=True)

        cols = st.columns(min(6, len(slides_rendered)))
        for i, img in enumerate(slides_rendered):
            with cols[i % len(cols)]:
                st.image(img, caption=f"Slide {i + 1}", use_container_width=True)

        col1, col2 = st.columns(2)
        with col1:
            zip_buf = io.BytesIO()
            with zipfile.ZipFile(zip_buf, "w", zipfile.ZIP_DEFLATED) as zf:
                for i, img in enumerate(slides_rendered):
                    b = io.BytesIO()
                    img.save(b, format="PNG")
                    zf.writestr(f"slide_{i + 1:02d}.png", b.getvalue())
            st.download_button("⬇️ Download All Slides (ZIP of PNGs)", data=zip_buf.getvalue(),
                                file_name="ai_deck_slides.zip", mime="application/zip", use_container_width=True)
        with col2:
            pptx_bytes = build_pptx(plan, theme, accent, logo_image, st.session_state.deck_images, logo_position)
            st.download_button("⬇️ Download Editable PowerPoint (.pptx)", data=pptx_bytes,
                                file_name="ai_deck.pptx",
                                mime="application/vnd.openxmlformats-officedocument.presentationml.presentation",
                                use_container_width=True)


if __name__ == "__main__":
    main()
